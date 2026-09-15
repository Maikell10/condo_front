import csv
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

XLSM = Path(r"c:\Dev\ProyectoCondominio\condominio-admin\RECIBOS CONDOMINIO  JUNIO 2026.xlsm")
OUT_IMPORT = Path(r"c:\Dev\ProyectoCondominio\condo-backend\istmo-import-complex-data.csv")
OUT_LEGACY = Path(r"c:\Dev\ProyectoCondominio\condo-backend\istmo-subir-conjunto-2.csv")
OUT_REPORT = Path(r"c:\Dev\ProyectoCondominio\condominio-admin\docs\context\istmo-import-report.txt")
DEFAULT_PWD = "$2b$10$P6nti1ZV4YagWwXoYnp3U.plcyRTPloHT.ylsvmlRwHM0y1LwJUiK"

SHEET_NAME_PATTERN = re.compile(r"^(.+)-(\d+)$", re.UNICODE)


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_sheet(workbook, name):
    worksheet = workbook[name]
    owner = ""
    header_row_index = None

    rows = list(worksheet.iter_rows(min_row=1, max_row=20, values_only=True))
    for index, row in enumerate(rows):
        cells = [clean(cell) for cell in row]
        if any("PROPIETARIO" in cell.upper() for cell in cells) and any(
            "ALICUOTA" in cell.upper() for cell in cells
        ):
            header_row_index = index
            break

    if header_row_index is not None:
        for row in rows[header_row_index + 1 : header_row_index + 4]:
            cells = [clean(cell) for cell in row]
            candidate_owner = cells[3] if len(cells) > 3 else ""

            if candidate_owner and candidate_owner not in ("0", "-", "."):
                owner = candidate_owner
                break

    match = SHEET_NAME_PATTERN.match(name.strip())
    if not match:
        return None

    building, apartment = match.group(1).upper(), match.group(2)
    owner = owner if owner and owner not in ("0", "-", ".") else ""

    return {
        "building_name": building,
        "apartment": apartment,
        "number": f"{building}-{apartment}",
        "owner": owner,
        "sheet": name,
    }


def assign_equal_alicuotas(records):
    by_building = defaultdict(list)
    for record in records:
        by_building[record["building_name"]].append(record)

    for building, building_records in by_building.items():
        equal_share = round(1 / len(building_records), 6)
        for record in building_records:
            record["alicuota"] = f"{equal_share:.6f}".rstrip("0").rstrip(".")


def building_sort_key(building_order, record):
    return (building_order.index(record["building_name"]), int(record["apartment"]))


def main():
    workbook = openpyxl.load_workbook(
        XLSM, read_only=True, data_only=True, keep_vba=True
    )

    building_order = []
    seen_buildings = set()
    for sheet_name in workbook.sheetnames:
        if sheet_name.lower() == "mes":
            continue
        match = SHEET_NAME_PATTERN.match(sheet_name.strip())
        if not match:
            continue
        building = match.group(1).upper()
        if building not in seen_buildings:
            seen_buildings.add(building)
            building_order.append(building)

    records = [
        parse_sheet(workbook, sheet_name)
        for sheet_name in workbook.sheetnames
        if sheet_name.lower() != "mes"
    ]
    records = [record for record in records if record]
    assign_equal_alicuotas(records)
    records.sort(key=lambda item: building_sort_key(building_order, item))
    workbook.close()

    with OUT_IMPORT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(
            ["building_name", "apartment", "Nombre Propietario", "email", "alicuota"]
        )
        for record in records:
            writer.writerow(
                [
                    record["building_name"],
                    record["apartment"],
                    record["owner"],
                    "",
                    record["alicuota"],
                ]
            )

    with OUT_LEGACY.open("w", encoding="latin-1", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(
            [
                "building_id",
                "building_name",
                "number",
                "alicuota",
                "name_user",
                "email",
                "password",
                "role",
                "deuda",
            ]
        )
        for record in records:
            writer.writerow(
                [
                    "",
                    record["building_name"],
                    record["number"],
                    record["alicuota"].replace(".", ","),
                    record["owner"],
                    "",
                    DEFAULT_PWD,
                    "OWNER",
                    "0,00",
                ]
            )

    no_owner = [record for record in records if not record["owner"]]
    by_building = defaultdict(list)
    for record in records:
        by_building[record["building_name"]].append(record)

    lines = [
        f"Total apartamentos: {len(records)}",
        f"Edificios: {len(building_order)}",
        f"Con propietario: {len(records) - len(no_owner)}",
        f"Sin propietario: {len(no_owner)}",
        f"Alicuota: partes iguales por edificio (1/n)",
        "",
        "Orden edificios:",
        f"  {', '.join(building_order)}",
        "",
        "Edificios y cantidad:",
    ]
    for building in building_order:
        share = by_building[building][0]["alicuota"]
        lines.append(f"  {building}: {len(by_building[building])} aptos | alicuota c/u {share}")
    lines.extend(["", "Apartamentos sin propietario:"])
    for record in no_owner:
        lines.append(
            f"  {record['sheet']} | edificio {record['building_name']} | apt {record['apartment']}"
        )

    OUT_REPORT.write_text("\n".join(lines), encoding="utf-8")

    print(f"Wrote {OUT_IMPORT}")
    print(f"Apartamentos: {len(records)} | Edificios: {len(building_order)}")
    print(f"Edificios: {building_order}")


if __name__ == "__main__":
    main()
